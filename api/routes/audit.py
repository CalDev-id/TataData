import io
import re
import threading
import time
from copy import copy, deepcopy
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, BackgroundTasks, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel

from features.audit_db.pipeline import RUNTIMES, pipeline_audit_db, prepare_audit_db, run_normalization


router = APIRouter()
jobs = {}
normalization_jobs = {}
jobs_lock = threading.RLock()
TEMPLATE = Path(__file__).resolve().parents[2] / "template" / "template.xlsx"
HEADERS = (
    "No", "Nama Database", "Nama Tabel", "Objek Temuan", "Aturan Playbook/UU",
    "Uraian Ketidaksesuaian", "Solusi", "Status", "Catatan", "Jenis Temuan",
)
FIELDS = (
    "no", "nama_database", "nama_tabel", "objek_temuan", "aturan_playbook",
    "uraian_ketidaksesuaian", "solusi", "status", "catatan", "jenis_temuan",
)


class NormalizationRequest(BaseModel):
    audit_job_id: str
    model: str


def _elapsed(job):
    return int(time.monotonic() - job["started_at"]) if job["started_at"] else 0


def _event(job, stage, kind, message):
    # Format stage/kind dipakai front end untuk menampilkan log progres.
    job["process_events"].append({
        "stage": stage, "kind": kind, "message": message,
        "elapsed_seconds": _elapsed(job),
    })


def _finish_if_terminal(job):
    if job["status"] in {"completed", "failed"} and job["normalization_status"] in {"completed", "failed"}:
        job["finished_at"] = time.monotonic()


def _get_job(job_id):
    with jobs_lock:
        now = time.monotonic()
        for old_id, old_job in list(jobs.items()):
            if old_job["finished_at"] and now - old_job["finished_at"] > 3600:
                normalization_jobs.pop(old_job["normalization_id"], None)
                del jobs[old_id]
        job = jobs.get(job_id)
        if job is None:
            raise HTTPException(404, "Proses tidak ditemukan atau sudah kedaluwarsa.")
        return job


def _update_job(job, change, **details):
    # Dua proses melapor bersamaan; lock menjaga status dan daftar temuan tetap konsisten.
    with jobs_lock:
        if change == "audit_batch_started":
            _event(job, "audit", "started", f"Audit batch {details['batch']}/{job['total_batches']} dimulai.")
        elif change == "audit_batch_completed":
            job["audit_findings"].extend(details["findings"])
            job["completed_batches"] = details["batch"]
            _event(job, "audit", "completed", f"Audit batch {details['batch']}/{job['total_batches']} selesai.")
        elif change == "audit_completed":
            job["status"] = "completed"
        elif change == "audit_batch_failed":
            job["status"] = "failed"
            job["audit_error"] = details["error"]
            _event(job, "audit", "failed", f"Audit batch {details['batch']}/{job['total_batches']} gagal: {details['error']}")
        elif change == "normalization_started":
            job["normalization_status"] = "running"
            _event(job, "normalization", "started", "Normalisasi dimulai.")
        elif change == "normalization_completed":
            job["normalization_findings"] = details["findings"]
            job["normalization_status"] = "completed"
            _event(job, "normalization", "completed", f"Normalisasi selesai: {len(details['findings'])} temuan.")
        elif change == "normalization_failed":
            job["normalization_status"] = "failed"
            job["normalization_error"] = details["error"]
            _event(job, "normalization", "failed", f"Normalisasi gagal: {details['error']}")
        _finish_if_terminal(job)


def _process_job(job_id):
    job = _get_job(job_id)
    with jobs_lock:
        job["started_at"] = job["started_at"] or time.monotonic()
        job["status"] = "running"
        job["finished_at"] = None

    def update_status(change, **details):
        _update_job(job, change, **details)

    try:
        pipeline_audit_db(
            job["prepared"], job["model"], update_status,
            job["completed_batches"], len(job["audit_findings"]),
            start_normalization=job["normalization_status"] == "queued",
        )
    except Exception as exc:
        with jobs_lock:
            if job["status"] == "completed":
                job["normalization_status"] = "failed"
                job["normalization_error"] = str(exc)
                _event(job, "normalization", "failed", f"Normalisasi gagal: {exc}")
            else:
                job["status"] = "failed"
                job["audit_error"] = str(exc)
                batch = job["completed_batches"] + 1
                _event(job, "audit", "failed", f"Audit batch {batch}/{job['total_batches']} gagal: {exc}")
            _finish_if_terminal(job)


def _retry_normalization_job(job_id):
    job = _get_job(job_id)

    def update_status(change, **details):
        _update_job(job, change, **details)

    run_normalization(
        job["prepared"]["ddl"], job["normalization_model"], update_status,
        job["prepared"]["database_name"],
    )


def _status(job):
    status_url = f"/api/v1/audit/spreadsheet/{job['job_id']}"
    result = {
        "job_id": job["job_id"], "status_url": status_url,
        "status": job["status"], "source_filename": job["source_filename"],
        "completed_batches": job["completed_batches"],
        "total_batches": job["total_batches"], "elapsed_seconds": _elapsed(job),
        "process_events": list(job["process_events"]),
        "retry_available": job["status"] == "failed",
        "error": job["audit_error"],
    }
    result["normalization_job_id"] = job["normalization_id"]
    result["normalization_status_url"] = f"/api/v1/normalization/{job['normalization_id']}"
    result["normalization_error"] = job["normalization_error"]
    return result


@router.post("/api/v1/audit/spreadsheet")
async def upload_audit(background_tasks: BackgroundTasks, file: UploadFile = File(...), model: str = Form(...)):
    if model not in RUNTIMES:
        raise HTTPException(422, "Model analisis tidak tersedia.")
    filename = Path(file.filename or "").name
    if Path(filename).suffix.lower() not in {".sql", ".ddl"}:
        raise HTTPException(422, "Unggah file .sql atau .ddl.")
    try:
        prepared = prepare_audit_db(await file.read())
    except UnicodeDecodeError as exc:
        raise HTTPException(422, "File harus menggunakan encoding UTF-8.") from exc
    if not prepared["table_count"]:
        raise HTTPException(422, "File harus berisi CREATE TABLE MySQL.")

    job_id = uuid4().hex
    normalization_id = uuid4().hex
    job = {
        "job_id": job_id, "normalization_id": normalization_id,
        "source_filename": filename, "prepared": prepared,
        "model": model, "normalization_model": model,
        "status": "queued", "normalization_status": "queued",
        "completed_batches": 0, "total_batches": len(prepared["batches"]),
        "audit_findings": [], "normalization_findings": [], "process_events": [],
        "audit_error": None, "normalization_error": None,
        "started_at": None, "finished_at": None,
    }
    with jobs_lock:
        jobs[job_id] = job
        normalization_jobs[normalization_id] = job_id
    background_tasks.add_task(_process_job, job_id)
    return _status(job)


@router.get("/api/v1/audit/spreadsheet/{job_id}")
def audit_status(job_id: str):
    job = _get_job(job_id)
    with jobs_lock:
        return _status(job)


@router.post("/api/v1/audit/spreadsheet/{job_id}/retry")
def retry_audit(job_id: str, background_tasks: BackgroundTasks):
    job = _get_job(job_id)
    with jobs_lock:
        if job["status"] != "failed":
            raise HTTPException(409, "Audit tidak dalam status gagal.")
        job["status"] = "queued"
        job["audit_error"] = None
        job["finished_at"] = None
        _event(job, "audit", "retry", f"Audit batch {job['completed_batches'] + 1}/{job['total_batches']}: mencoba lagi.")
        status = _status(job)
    background_tasks.add_task(_process_job, job_id)
    return status


@router.get("/api/v1/normalization/{normalization_id}")
def normalization_status(normalization_id: str):
    with jobs_lock:
        job_id = normalization_jobs.get(normalization_id)
    if not job_id:
        raise HTTPException(404, "Proses normalisasi tidak ditemukan.")
    job = _get_job(job_id)
    with jobs_lock:
        return {
            "job_id": normalization_id, "audit_job_id": job_id,
            "status": job["normalization_status"], "error": job["normalization_error"],
            "process_events": list(job["process_events"]),
            "findings_url": f"/api/v1/audit/spreadsheet/{job_id}/combined-findings"
            if job["status"] == "completed" and job["normalization_status"] == "completed" else None,
        }


@router.post("/api/v1/normalization")
def retry_normalization(payload: NormalizationRequest, background_tasks: BackgroundTasks):
    job = _get_job(payload.audit_job_id)
    if payload.model not in RUNTIMES:
        raise HTTPException(422, "Model analisis tidak tersedia.")
    with jobs_lock:
        if job["normalization_status"] != "failed":
            raise HTTPException(409, "Normalisasi tidak dalam status gagal.")
        normalization_jobs.pop(job["normalization_id"], None)
        normalization_id = uuid4().hex
        job["normalization_id"] = normalization_id
        job["normalization_model"] = payload.model
        job["normalization_status"] = "queued"
        job["normalization_error"] = None
        job["finished_at"] = None
        normalization_jobs[normalization_id] = job["job_id"]
        _event(job, "normalization", "retry", "Mengulang Normalisasi.")
    background_tasks.add_task(_retry_normalization_job, job["job_id"])
    return {
        "job_id": normalization_id,
        "status_url": f"/api/v1/normalization/{normalization_id}",
    }


def _completed_job(job_id):
    job = _get_job(job_id)
    if job["status"] != "completed" or job["normalization_status"] != "completed":
        raise HTTPException(409, "Hasil Audit dan Normalisasi belum siap.")
    return job


@router.get("/api/v1/audit/spreadsheet/{job_id}/combined-findings")
def combined_findings(job_id: str):
    job = _completed_job(job_id)
    with jobs_lock:
        return {
            "summary": {
                "table_count": job["prepared"]["table_count"],
                "column_count": job["prepared"]["column_count"],
            },
            "findings": job["audit_findings"] + job["normalization_findings"],
        }


def _fill_sheet(sheet, findings):
    sheet["B6"] = len(findings)
    sheet["E6"] = sum(row["status"] == "To Do" for row in findings)
    sheet["G6"] = len(findings) - sheet["E6"].value
    sheet["I6"] = sheet["G6"].value / len(findings) if findings else 0
    sheet["I6"].number_format = "0%"
    for column, header in enumerate(HEADERS, 2):
        cell = sheet.cell(8, column)
        cell.value = header
        if column == 11:
            cell._style = copy(sheet["J8"]._style)
    sheet.column_dimensions["K"].width = 20
    sheet.freeze_panes = "A9"
    sheet.auto_filter.ref = f"B8:K{max(8, 8 + len(findings))}"
    for row in sheet.iter_rows(min_row=9, max_col=11):
        for cell in row[1:]:
            cell.value = None
    for row_number, finding in enumerate(findings, 9):
        for column, field in enumerate(FIELDS, 2):
            cell = sheet.cell(row_number, column)
            if row_number > 9:
                cell._style = copy(sheet.cell(9, column)._style)
            value = finding[field]
            if field == "no":
                cell.value = value
            else:
                cell.value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
                cell.data_type = "s"


@router.get("/api/v1/audit/spreadsheet/{job_id}/download/combined")
def download_combined(job_id: str):
    job = _completed_job(job_id)
    with jobs_lock:
        audit_findings = list(job["audit_findings"])
        normalization_findings = list(job["normalization_findings"])
    workbook = load_workbook(TEMPLATE)
    audit_sheet = workbook.worksheets[0]
    normalization_sheet = workbook.copy_worksheet(audit_sheet)
    normalization_sheet.conditional_formatting = deepcopy(audit_sheet.conditional_formatting)
    for sheet in list(workbook.worksheets):
        if sheet not in (audit_sheet, normalization_sheet):
            workbook.remove(sheet)
    audit_sheet.title = "Audit Playbook"
    normalization_sheet.title = "Normalisasi"
    _fill_sheet(audit_sheet, audit_findings)
    _fill_sheet(normalization_sheet, normalization_findings)
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="temuan_audit_normalisasi.xlsx"'},
    )
